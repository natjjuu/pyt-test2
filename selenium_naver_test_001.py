from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
import time
from openpyxl import load_workbook

# -----------------------------
# 엑셀 로드
# -----------------------------
excel_path = "testcase1.xlsx"
wb = load_workbook(excel_path)
ws = wb.active

# -----------------------------
# 결과 기록 함수
# -----------------------------
def write_result(cell, result):
    ws[cell] = result

# -----------------------------
# 테스트 시작
# -----------------------------
try:
    # 1. 크롬 실행
    driver = webdriver.Chrome()
    write_result("C8", "PASS")
except Exception as e:
    write_result("C8", "FAIL")
    wb.save(excel_path)
    print("크롬 실행 오류:", e)
    exit()

time.sleep(3)

try:
    # 2. 네이버 접속
    driver.get("https://www.naver.com")
    time.sleep(3)

    # 접속 여부 확인 (title에 'NAVER' 포함 확인)
    if "NAVER" in driver.title:
        write_result("C9", "PASS")
    else:
        write_result("C9", "FAIL")
except Exception as e:
    write_result("C9", "FAIL")
    wb.save(excel_path)
    print("네이버 접속 오류:", e)
    driver.quit()
    exit()

try:
    # 3. 날씨 검색
    search = driver.find_element(By.ID, "query")
    search.send_keys("날씨")
    search.send_keys(Keys.ENTER)
    time.sleep(3)

    # 날씨 페이지 여부 확인
    if "날씨" in driver.title:
        write_result("C10", "PASS")
    else:
        write_result("C10", "FAIL")
except Exception as e:
    write_result("C10", "FAIL")
    wb.save(excel_path)
    print("날씨 검색 오류:", e)
    driver.quit()
    exit()

try:
    # 4. 뒤로 가기 → 네이버로 복귀
    driver.back()
    time.sleep(3)

    if "NAVER" in driver.title:
        write_result("C11", "PASS")
    else:
        write_result("C11", "FAIL")
except Exception as e:
    write_result("C11", "FAIL")
    wb.save(excel_path)
    print("뒤로가기 오류:", e)
    driver.quit()
    exit()

try:
    # 5. 크롬 종료
    driver.quit()
    write_result("C12", "PASS")
except Exception as e:
    write_result("C12", "FAIL")
    print("브라우저 종료 오류:", e)

# -----------------------------
# 엑셀 저장
# -----------------------------
wb.save(excel_path)
print("테스트 완료! 결과는 testcase1.xlsx에 저장되었습니다.")
